from datetime import date
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from docx import Document
from pypdf import PdfReader
import pytest

from app.domain import ServiceItem
from app.pricing.catalog import ServiceCatalog
from app.pricing.engine import PricingEngine, VolumeDiscountRule
from app.pricing.models import RequestedService
from app.proposals.errors import PDFConversionError, PricingIntegrityError
from app.proposals.generator import ProposalGenerator, format_money
from app.proposals.models import EvidenceFact, ProposalDocumentData


def _data() -> ProposalDocumentData:
    catalog = ServiceCatalog(items={
        "AUTO_SETUP": ServiceItem(service_code="AUTO_SETUP", service_name="Customer request automation", unit="project", unit_price=Decimal("1500.00"), minimum_quantity=1, active=True, currency="EUR"),
        "TRAINING": ServiceItem(service_code="TRAINING", service_name="Staff training", unit="session", unit_price=Decimal("250.00"), minimum_quantity=1, active=True, currency="EUR"),
    }, currency="EUR")
    pricing = PricingEngine(catalog, discount_rule=VolumeDiscountRule()).calculate([
        RequestedService(service_code="AUTO_SETUP", quantity=2),
        RequestedService(service_code="TRAINING", quantity=2),
    ])
    return ProposalDocumentData(
        proposal_id=uuid4(), prepared_on=date(2026, 9, 28), valid_until=date(2026, 10, 12),
        client_company="Example Manufacturing Ltd", client_contact="Alex Morgan",
        project_title="Customer Request Automation Implementation",
        executive_summary="Northstar Automation proposes a controlled customer request workflow with structured intake, internal notifications and staff training.",
        scope_items=["Configure customer request intake and validation.", "Connect approved Google Sheets and email notification steps.", "Provide two staff training sessions."],
        evidence=[
            EvidenceFact(statement="Customer request automation includes email notifications and Google Sheets integration.", source_id="catalogue-page-1", source_location="catalogue.pdf - page 1"),
            EvidenceFact(statement="Staff training is available as an optional service.", source_id="terms-page-2", source_location="terms.docx - paragraph 4"),
        ], pricing=pricing,
    )


def _docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def test_generates_docx_pdf_with_matching_money(tmp_path: Path) -> None:
    data = _data(); docx = tmp_path / "proposal.docx"; pdf = tmp_path / "proposal.pdf"
    ProposalGenerator().generate(data, docx, pdf)
    docx_text = _docx_text(docx); pdf_text = _pdf_text(pdf)
    expected = {format_money(data.pricing.currency, value) for value in [data.pricing.subtotal, -data.pricing.discount, data.pricing.taxable_amount, data.pricing.tax, data.pricing.total]}
    for amount in expected:
        assert amount in docx_text
        assert amount in pdf_text
    assert "DRAFT - PENDING MANAGER APPROVAL" in docx_text
    assert "DRAFT - PENDING MANAGER APPROVAL" in pdf_text
    assert "catalogue-page-1" in docx_text and "catalogue-page-1" in pdf_text


def test_rejects_corrupted_pricing_totals(tmp_path: Path) -> None:
    data = _data()
    corrupted = data.model_copy(update={"pricing": data.pricing.model_copy(update={"total": Decimal("1.00")})})
    with pytest.raises(PricingIntegrityError):
        ProposalGenerator().create_docx(corrupted, tmp_path / "bad.docx")


def test_pdf_conversion_reports_missing_converter(monkeypatch, tmp_path: Path) -> None:
    docx = tmp_path / "proposal.docx"
    ProposalGenerator().create_docx(_data(), docx)
    monkeypatch.setattr("app.proposals.generator.shutil.which", lambda _: None)
    monkeypatch.delenv("PROGRAMFILES", raising=False)
    monkeypatch.delenv("PROGRAMFILES(X86)", raising=False)
    with pytest.raises(PDFConversionError, match="not found"):
        ProposalGenerator().convert_to_pdf(docx, tmp_path / "proposal.pdf")


def test_rejects_damaged_docx_conversion(tmp_path: Path) -> None:
    damaged = tmp_path / "damaged.docx"; damaged.write_bytes(b"not a document")
    with pytest.raises(PDFConversionError):
        ProposalGenerator().convert_to_pdf(damaged, tmp_path / "damaged.pdf")
