from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.ingestion import DocumentIngestionService


def _write_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 750, "Northstar Automation service catalogue")
    pdf.drawString(72, 730, "Customer request automation includes email notifications.")
    pdf.save()


def _write_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Northstar Automation Terms", level=1)
    document.add_paragraph("Delivery starts after manager approval and requirements review.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Support"
    table.cell(1, 1).text = "30 days"
    document.save(path)


def _write_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Price List"
    sheet.append(["service_code", "service_name", "unit_price", "currency"])
    sheet.append(["AUTO_SETUP", "Automation setup", 750, "EUR"])
    workbook.save(path)


def test_ingests_pdf_with_page_metadata(tmp_path: Path) -> None:
    path = tmp_path / "catalogue.pdf"
    _write_pdf(path)

    result = DocumentIngestionService(tmp_path).ingest(path)

    assert result.metadata.extension == ".pdf"
    assert result.metadata.parser == "pypdf"
    assert result.segments[0].location == "page 1"
    assert "Northstar Automation" in result.chunks[0].text
    assert result.chunks[0].metadata["filename"] == "catalogue.pdf"
    assert result.chunks[0].metadata["extension"] == ".pdf"


def test_ingests_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "terms.docx"
    _write_docx(path)

    result = DocumentIngestionService(tmp_path).ingest(path)

    locations = {segment.location for segment in result.segments}
    assert result.metadata.parser == "python-docx"
    assert "paragraph 2" in locations
    assert "table 1, row 2" in locations
    assert any("30 days" in chunk.text for chunk in result.chunks)


def test_ingests_xlsx_rows_with_sheet_metadata(tmp_path: Path) -> None:
    path = tmp_path / "prices.xlsx"
    _write_xlsx(path)

    result = DocumentIngestionService(tmp_path).ingest(path)

    assert result.metadata.parser == "openpyxl"
    assert result.segments[1].metadata == {"sheet": "Price List", "row": "2"}
    assert "AUTO_SETUP" in result.segments[1].text


def test_ingests_utf8_csv_and_normalizes_whitespace(tmp_path: Path) -> None:
    path = tmp_path / "addons.csv"
    path.write_text(
        "service_code,description,price\nTRAINING,  Staff   training  ,250\n",
        encoding="utf-8",
    )

    result = DocumentIngestionService(tmp_path).ingest(path)

    assert result.metadata.parser == "python-csv"
    assert result.segments[1].location == "row 2"
    assert result.segments[1].text == "TRAINING | Staff training | 250"


def test_same_file_produces_stable_document_and_chunk_ids(tmp_path: Path) -> None:
    path = tmp_path / "services.csv"
    path.write_text("code,description\nA1,Automation service\n", encoding="utf-8")
    service = DocumentIngestionService(tmp_path)

    first = service.ingest(path)
    second = service.ingest(path)

    assert first.metadata.document_id == second.metadata.document_id
    assert [chunk.chunk_id for chunk in first.chunks] == [chunk.chunk_id for chunk in second.chunks]
