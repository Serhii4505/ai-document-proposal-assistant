"""Bounded parsers for supported business document formats."""

from __future__ import annotations

import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.ingestion.errors import CorruptDocumentError, FileLimitError, UnsupportedFileError
from app.ingestion.models import SourceSegment
from app.ingestion.normalize import normalize_text
from app.ingestion.security import IngestionLimits, ValidatedFile


def _segment(text: str, location: str, **metadata: str) -> SourceSegment | None:
    normalized = normalize_text(text)
    if not normalized:
        return None
    return SourceSegment(text=normalized, location=location, metadata=metadata)


def _safe_cell(value: object, limits: IngestionLimits) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        rendered = value.isoformat()
    elif isinstance(value, (int, float, Decimal, bool)):
        rendered = str(value)
    else:
        rendered = str(value)
    if len(rendered) > limits.max_cell_characters:
        raise FileLimitError("Spreadsheet or CSV cell is too large")
    return normalize_text(rendered)


def parse_pdf(file: ValidatedFile, limits: IngestionLimits) -> list[SourceSegment]:
    try:
        reader = PdfReader(file.path)
        if reader.is_encrypted:
            raise UnsupportedFileError("Encrypted PDF files are not supported")
        if len(reader.pages) > limits.max_pdf_pages:
            raise FileLimitError("PDF contains too many pages")
        segments = []
        for page_number, page in enumerate(reader.pages, start=1):
            item = _segment(page.extract_text() or "", f"page {page_number}", page=str(page_number))
            if item:
                segments.append(item)
        return segments
    except (FileLimitError, UnsupportedFileError):
        raise
    except Exception as exc:
        raise CorruptDocumentError("PDF could not be parsed") from exc


def parse_docx(file: ValidatedFile, limits: IngestionLimits) -> list[SourceSegment]:
    del limits
    try:
        document = Document(file.path)
        segments: list[SourceSegment] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            item = _segment(paragraph.text, f"paragraph {index}", paragraph=str(index))
            if item:
                segments.append(item)
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                text = " | ".join(normalize_text(cell.text) for cell in row.cells)
                item = _segment(
                    text,
                    f"table {table_number}, row {row_number}",
                    table=str(table_number),
                    row=str(row_number),
                )
                if item:
                    segments.append(item)
        return segments
    except Exception as exc:
        raise CorruptDocumentError("DOCX could not be parsed") from exc


def parse_xlsx(file: ValidatedFile, limits: IngestionLimits) -> list[SourceSegment]:
    try:
        workbook = load_workbook(file.path, read_only=True, data_only=True, keep_links=False)
        if len(workbook.worksheets) > limits.max_sheets:
            raise FileLimitError("Workbook contains too many sheets")
        segments: list[SourceSegment] = []
        for sheet in workbook.worksheets:
            safe_title = normalize_text(sheet.title)[:100]
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_number > limits.max_rows_per_sheet:
                    raise FileLimitError("Worksheet contains too many rows")
                if len(row) > limits.max_columns_per_row:
                    raise FileLimitError("Worksheet row contains too many columns")
                values = [_safe_cell(value, limits) for value in row]
                if not any(values):
                    continue
                item = _segment(
                    " | ".join(values),
                    f"sheet {safe_title}, row {row_number}",
                    sheet=safe_title,
                    row=str(row_number),
                )
                if item:
                    segments.append(item)
        workbook.close()
        return segments
    except FileLimitError:
        raise
    except Exception as exc:
        raise CorruptDocumentError("XLSX could not be parsed") from exc


def parse_csv(file: ValidatedFile, limits: IngestionLimits) -> list[SourceSegment]:
    try:
        with file.path.open("r", encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(8192)
            if "\x00" in sample:
                raise UnsupportedFileError("CSV file contains binary data")
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel

            reader = csv.reader(stream, dialect)
            segments: list[SourceSegment] = []
            for row_number, row in enumerate(reader, start=1):
                if row_number > limits.max_rows_per_sheet:
                    raise FileLimitError("CSV contains too many rows")
                if len(row) > limits.max_columns_per_row:
                    raise FileLimitError("CSV row contains too many columns")
                values = [_safe_cell(value, limits) for value in row]
                if not any(values):
                    continue
                item = _segment(
                    " | ".join(values),
                    f"row {row_number}",
                    row=str(row_number),
                )
                if item:
                    segments.append(item)
        return segments
    except (FileLimitError, UnsupportedFileError):
        raise
    except UnicodeDecodeError as exc:
        raise UnsupportedFileError("CSV must use UTF-8 encoding") from exc
    except (csv.Error, OSError) as exc:
        raise CorruptDocumentError("CSV could not be parsed") from exc


PARSERS = {
    ".pdf": ("pypdf", parse_pdf),
    ".docx": ("python-docx", parse_docx),
    ".xlsx": ("openpyxl", parse_xlsx),
    ".csv": ("python-csv", parse_csv),
}


def parse_document(file: ValidatedFile, limits: IngestionLimits) -> tuple[str, list[SourceSegment]]:
    parser_name, parser = PARSERS[file.extension]
    return parser_name, parser(file, limits)

