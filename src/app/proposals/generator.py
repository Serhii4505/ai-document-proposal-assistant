"""Styled DOCX generation, PDF conversion and pricing integrity checks."""

from __future__ import annotations

from decimal import Decimal
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader

from app.pricing.engine import money
from app.proposals.errors import PDFConversionError, PricingIntegrityError
from app.proposals.models import ProposalDocumentData


BLUE = "17365D"
LIGHT_BLUE = "DCE6F1"
GREY = "666666"


def format_money(currency: str, value: Decimal) -> str:
    return f"{currency} {money(value):,.2f}"


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


class ProposalGenerator:
    @staticmethod
    def validate_pricing(data: ProposalDocumentData) -> None:
        pricing = data.pricing
        subtotal = money(sum((line.line_total for line in pricing.line_items), Decimal("0")))
        if subtotal != money(pricing.subtotal):
            raise PricingIntegrityError("Subtotal does not match line items")
        taxable = money(pricing.subtotal - pricing.discount)
        tax = money(taxable * pricing.vat_rate)
        total = money(taxable + tax)
        if taxable != money(pricing.taxable_amount) or tax != money(pricing.tax) or total != money(pricing.total):
            raise PricingIntegrityError("Pricing totals are internally inconsistent")
        if pricing.discount < 0 or pricing.discount > pricing.subtotal:
            raise PricingIntegrityError("Discount is outside the allowed range")

    def create_docx(self, data: ProposalDocumentData, output_path: Path) -> Path:
        self.validate_pricing(data)
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75); section.right_margin = Inches(0.75)

        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = "Aptos"; normal.font.size = Pt(10)
        title_style = styles["Title"]
        title_style.font.name = "Aptos Display"; title_style.font.size = Pt(27)
        title_style.font.color.rgb = RGBColor.from_string(BLUE)
        for name, size in (("Heading 1", 16), ("Heading 2", 12)):
            style = styles[name]; style.font.name = "Aptos Display"; style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(BLUE)

        title = document.add_paragraph(style="Title")
        title.add_run("Commercial Proposal")
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = subtitle.add_run(data.project_title); run.bold = True; run.font.size = Pt(14)
        status = document.add_paragraph()
        status_run = status.add_run("DRAFT - PENDING MANAGER APPROVAL")
        status_run.bold = True; status_run.font.color.rgb = RGBColor(180, 60, 45)

        info = document.add_table(rows=4, cols=2)
        info.alignment = WD_TABLE_ALIGNMENT.LEFT
        info.style = "Table Grid"
        details = [
            ("Prepared for", f"{data.client_contact}, {data.client_company}"),
            ("Prepared by", "Northstar Automation"),
            ("Proposal ID", str(data.proposal_id)),
            ("Validity", f"{data.prepared_on.isoformat()} to {data.valid_until.isoformat()}"),
        ]
        for row, (label, value) in zip(info.rows, details, strict=True):
            row.cells[0].text = label; row.cells[1].text = value
            _shade(row.cells[0], LIGHT_BLUE); row.cells[0].paragraphs[0].runs[0].bold = True

        document.add_heading("Executive Summary", level=1)
        document.add_paragraph(data.executive_summary)
        document.add_heading("Proposed Scope", level=1)
        for item in data.scope_items:
            document.add_paragraph(item, style="List Bullet")

        document.add_heading("Confirmed Basis", level=1)
        document.add_paragraph("The following statements are based on approved source material.")
        evidence_table = document.add_table(rows=1, cols=2); evidence_table.style = "Table Grid"
        evidence_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, label in zip(evidence_table.rows[0].cells, ("Confirmed fact", "Source"), strict=True):
            cell.text = label; _shade(cell, BLUE); cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255); cell.paragraphs[0].runs[0].bold = True
        for fact in data.evidence:
            cells = evidence_table.add_row().cells
            cells[0].text = fact.statement
            cells[1].text = f"{fact.source_id} - {fact.source_location}"

        document.add_heading("Investment", level=1)
        price_table = document.add_table(rows=1, cols=5); price_table.style = "Table Grid"
        price_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headings = ("Service", "Unit", "Quantity", "Unit price", "Line total")
        for cell, label in zip(price_table.rows[0].cells, headings, strict=True):
            cell.text = label; _shade(cell, BLUE); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            run = cell.paragraphs[0].runs[0]; run.font.color.rgb = RGBColor(255, 255, 255); run.bold = True
        for line in data.pricing.line_items:
            cells = price_table.add_row().cells
            values = (f"{line.service_name}\n{line.service_code}", line.unit, str(line.quantity), format_money(data.pricing.currency, line.unit_price), format_money(data.pricing.currency, line.line_total))
            for cell, value in zip(cells, values, strict=True): cell.text = value

        totals = [
            ("Subtotal", data.pricing.subtotal),
            ("Discount", -data.pricing.discount),
            ("Taxable amount", data.pricing.taxable_amount),
            (f"VAT {data.pricing.vat_rate * 100:.0f}%", data.pricing.tax),
            ("Total", data.pricing.total),
        ]
        totals_table = document.add_table(rows=0, cols=2); totals_table.style = "Table Grid"; totals_table.alignment = WD_TABLE_ALIGNMENT.RIGHT
        for label, value in totals:
            cells = totals_table.add_row().cells; cells[0].text = label; cells[1].text = format_money(data.pricing.currency, value)
            if label == "Total":
                _shade(cells[0], LIGHT_BLUE); _shade(cells[1], LIGHT_BLUE)
                cells[0].paragraphs[0].runs[0].bold = True; cells[1].paragraphs[0].runs[0].bold = True

        document.add_page_break()
        document.add_heading("Terms and Approval", level=1)
        review = document.add_table(rows=4, cols=2); review.style = "Table Grid"
        review_details = [
            ("Proposal ID", str(data.proposal_id)),
            ("Client", data.client_company),
            ("Authoritative total", format_money(data.pricing.currency, data.pricing.total)),
            ("Valid until", data.valid_until.isoformat()),
        ]
        for row, (label, value) in zip(review.rows, review_details, strict=True):
            row.cells[0].text = label; row.cells[1].text = value
            _shade(row.cells[0], LIGHT_BLUE); row.cells[0].paragraphs[0].runs[0].bold = True
        document.add_paragraph()
        document.add_paragraph("This proposal is a draft. Scope, pricing and delivery may proceed only after manager approval. No service is authorized by this document alone.")
        document.add_paragraph("Manager decision:  [ ] Approve    [ ] Reject")
        document.add_paragraph("Name and signature: ______________________________    Date: ______________")

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run("Northstar Automation | Confidential proposal draft")
        footer_run.font.size = Pt(8); footer_run.font.color.rgb = RGBColor.from_string(GREY)
        document.save(output_path)
        return output_path

    def convert_to_pdf(self, docx_path: Path, pdf_path: Path) -> Path:
        windows_roots = [
            os.environ.get("PROGRAMFILES"),
            os.environ.get("PROGRAMFILES(X86)"),
        ]
        windows_candidates = [
            Path(root) / "LibreOffice" / "program" / "soffice.exe"
            for root in windows_roots
            if root
        ]
        executable = (
            shutil.which("soffice")
            or shutil.which("libreoffice")
            or next((str(path) for path in windows_candidates if path.is_file()), None)
        )
        if not executable:
            raise PDFConversionError("LibreOffice executable was not found")
        docx_path = Path(docx_path).resolve(); pdf_path = Path(pdf_path).resolve()
        try:
            with zipfile.ZipFile(docx_path) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise PDFConversionError("Input DOCX structure is invalid")
        except (zipfile.BadZipFile, OSError) as exc:
            raise PDFConversionError("Input DOCX is damaged or unreadable") from exc
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="proposal_lo_") as profile:
            result = subprocess.run([executable, "--headless", f"-env:UserInstallation={Path(profile).resolve().as_uri()}", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)], capture_output=True, text=True, timeout=60, check=False)
        generated = pdf_path.parent / f"{docx_path.stem}.pdf"
        if result.returncode != 0 or not generated.exists() or generated.stat().st_size == 0:
            raise PDFConversionError("DOCX to PDF conversion failed")
        if generated != pdf_path:
            generated.replace(pdf_path)
        try:
            reader = PdfReader(pdf_path)
            if not reader.pages:
                raise PDFConversionError("Converted PDF has no pages")
        except Exception as exc:
            if isinstance(exc, PDFConversionError): raise
            raise PDFConversionError("Converted PDF is invalid") from exc
        return pdf_path

    def generate(self, data: ProposalDocumentData, docx_path: Path, pdf_path: Path) -> tuple[Path, Path]:
        self.create_docx(data, docx_path)
        self.convert_to_pdf(docx_path, pdf_path)
        return Path(docx_path), Path(pdf_path)
