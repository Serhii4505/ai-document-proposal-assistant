"""Strict import of approved CSV/XLSX service catalogues."""

import csv
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from app.domain import ServiceItem
from app.ingestion.errors import UnsupportedFileError
from app.ingestion.security import IngestionLimits, validate_file
from app.pricing.errors import CatalogError


REQUIRED_COLUMNS = {"service_code", "service_name", "unit", "unit_price", "minimum_quantity", "active", "currency"}


@dataclass(frozen=True, slots=True)
class ServiceCatalog:
    items: dict[str, ServiceItem]
    currency: str

    def get(self, service_code: str) -> ServiceItem | None:
        return self.items.get(service_code)


def _parse_active(value: object) -> bool:
    normalized = str(value).strip().casefold()
    if normalized in {"true", "1", "yes"}:
        return True
    if normalized in {"false", "0", "no"}:
        return False
    raise CatalogError("active must be true or false")


def _service_item(row: dict[str, object], row_number: int) -> ServiceItem:
    missing = REQUIRED_COLUMNS - row.keys()
    if missing:
        raise CatalogError(f"Catalogue is missing columns: {', '.join(sorted(missing))}")
    try:
        return ServiceItem(
            service_code=str(row["service_code"]).strip().upper(),
            service_name=str(row["service_name"]).strip(),
            unit=str(row["unit"]).strip(),
            unit_price=Decimal(str(row["unit_price"]).strip()),
            minimum_quantity=Decimal(str(row["minimum_quantity"]).strip()),
            active=_parse_active(row["active"]),
            currency=str(row["currency"]).strip().upper(),
        )
    except (InvalidOperation, ValidationError, CatalogError) as exc:
        raise CatalogError(f"Invalid catalogue row {row_number}") from exc


def _read_csv(path: Path, limits: IngestionLimits) -> list[tuple[int, dict[str, object]]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if reader.fieldnames is None:
                raise CatalogError("Catalogue header is missing")
            rows = []
            for number, row in enumerate(reader, start=2):
                if number > limits.max_rows_per_sheet + 1:
                    raise CatalogError("Catalogue contains too many rows")
                rows.append((number, dict(row)))
            return rows
    except UnicodeDecodeError as exc:
        raise CatalogError("CSV catalogue must use UTF-8") from exc


def _read_xlsx(path: Path, limits: IngestionLimits) -> list[tuple[int, dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise CatalogError("Price catalogue must contain exactly one worksheet")
        iterator = workbook.active.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            raise CatalogError("Catalogue header is missing")
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        rows = []
        for number, values in enumerate(iterator, start=2):
            if number > limits.max_rows_per_sheet + 1:
                raise CatalogError("Catalogue contains too many rows")
            if any(value is not None and str(value).strip() for value in values):
                rows.append((number, dict(zip(headers, values, strict=False))))
        return rows
    finally:
        workbook.close()


def load_catalog(path: Path, *, allowed_root: Path, limits: IngestionLimits | None = None) -> ServiceCatalog:
    active_limits = limits or IngestionLimits()
    validated = validate_file(path, allowed_root=allowed_root, limits=active_limits)
    if validated.extension not in {".csv", ".xlsx"}:
        raise UnsupportedFileError("Price catalogue must be CSV or XLSX")
    rows = _read_csv(validated.path, active_limits) if validated.extension == ".csv" else _read_xlsx(validated.path, active_limits)
    if not rows:
        raise CatalogError("Catalogue contains no service rows")

    items: dict[str, ServiceItem] = {}
    currencies: set[str] = set()
    for number, row in rows:
        item = _service_item(row, number)
        if item.service_code in items:
            raise CatalogError(f"Duplicate service_code: {item.service_code}")
        items[item.service_code] = item
        currencies.add(item.currency)
    if len(currencies) != 1:
        raise CatalogError("All catalogue rows must use the same currency")
    return ServiceCatalog(items=items, currency=currencies.pop())

